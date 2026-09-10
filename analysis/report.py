"""Analiz sonuclarindan cok sayfali Excel raporu uretir."""
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CATALOG = ROOT / "analysis" / "data" / "catalog.json"
GUTENBERG = ROOT / "public" / "analysis" / "gutenberg.json"
MIGRATION = ROOT / "public" / "analysis" / "migration.json"
OUT = ROOT / "analysis" / "output" / "afet-paneli-analiz-raporu.xlsx"

FONT = "Arial"
NAVY = "1E293B"
ACCENT = "0F766E"
LIGHT = "F1F5F9"
WARN = "FEF3C7"

thin = Side(style="thin", color="CBD5E1")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, last_col):
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=NAVY)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = Font(name=FONT, size=10, italic=True, color="64748B")
    ws.row_dimensions[1].height = 24


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def body(ws, first_row, last_row, last_col, zebra=True):
    for r in range(first_row, last_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            if zebra and (r - first_row) % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=LIGHT)


# ---------------------------------------------------------------- sayfa 1
def sheet_ozet(wb, cat_stats, n_countries, n_faults):
    ws = wb.active
    ws.title = "Özet"
    widths(ws, {"A": 38, "B": 22, "C": 46})

    title(ws, "Küresel Doğal Afet İzleme Paneli",
          f"Analiz raporu · {datetime.now():%d.%m.%Y %H:%M}")

    rows = [
        ("VERİ KAYNAĞI", "", ""),
        ("Katalog", "USGS FDSN", "Kamuya açık, anahtar gerektirmez"),
        ("Kapsam", f"{cat_stats['first_year']}–{cat_stats['last_year']}", "Aletsel dönem"),
        ("Büyüklük eşiği", "M4.5+", "Alt eşikte sensör yoğunluğu haritayı çarpıtıyor"),
        ("Toplam deprem", cat_stats["total"], "Ham katalog kaydı"),
        ("", "", ""),
        ("ANALİZLER", "", ""),
        ("Gutenberg-Richter", f"{n_countries} bölge", "log₁₀(N) = a − b·M"),
        ("Göç analizi", f"{n_faults} fay hattı", "Zaman-mesafe regresyonu"),
        ("Anlamlılık testi", "Permütasyon", "5.000 rastgele karıştırma"),
        ("Çoklu karşılaştırma", "Bonferroni", "α = 0,05 / test sayısı"),
        ("", "", ""),
        ("SINIRLILIKLAR", "", ""),
        ("Tarihsel eksiklik", "1900–1960", "Küçük depremler kayda geçmemiş"),
        ("Konum hassasiyeti", "±25 km", "Eski kayıtlarda daha yüksek belirsizlik"),
        ("Deprem tahmini", "Yapılmaz", "Bilimsel olarak mümkün değildir"),
    ]

    r = 4
    for label, value, note in rows:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=value)
        ws.cell(row=r, column=3, value=note)

        is_section = label and not value and not note
        ws.cell(row=r, column=1).font = Font(
            name=FONT, size=10, bold=bool(label), color=ACCENT if is_section else "000000"
        )
        ws.cell(row=r, column=2).font = Font(name=FONT, size=10)
        ws.cell(row=r, column=3).font = Font(name=FONT, size=9, color="64748B")

        if is_section:
            for c in range(1, 4):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=LIGHT)
        r += 1

    ws.cell(row=8, column=2).number_format = "#,##0"

    ws.cell(row=r + 1, column=1, value="Bu rapor analiz betiklerinden otomatik üretilmiştir.")
    ws.cell(row=r + 1, column=1).font = Font(name=FONT, size=9, italic=True, color="94A3B8")
    return ws


# ---------------------------------------------------------------- sayfa 2
def sheet_gutenberg(wb, gr):
    ws = wb.create_sheet("Gutenberg-Richter")
    widths(ws, {"A": 26, "B": 11, "C": 11, "D": 11, "E": 13, "F": 13, "G": 13, "H": 24})

    title(ws, "Gutenberg-Richter b Değerleri",
          "Düşük b: kilitli fay, büyük olaylarla boşalma · Yüksek b: küçük olaylarla sızma")

    headers = ["Bölge", "b değeri", "a değeri", "R²", "Olay sayısı",
               "Katalog (yıl)", "M7 tekrar (yıl)", "Değerlendirme"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))

    items = sorted(
        ((k, v) for k, v in gr.items() if k != "global"),
        key=lambda kv: kv[1]["b"],
    )
    if "global" in gr:
        items.insert(0, ("DÜNYA", gr["global"]))

    r = 5
    for name, f in items:
        rec = f.get("recurrence", {}).get("M7.0")
        if f["n"] < 500 and rec and rec > 500:
            rec = None

        if f["b"] < 0.85:
            note = "Kilitli fay eğilimi"
        elif f["b"] > 1.05:
            note = "Volkanik / gerilmeli"
        else:
            note = "Küresel ortalamaya yakın"

            
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=f["b"]).number_format = "0.000"
        ws.cell(row=r, column=3, value=f["a"]).number_format = "0.00"
        ws.cell(row=r, column=4, value=f["r2"]).number_format = "0.0000"
        ws.cell(row=r, column=5, value=f["n"]).number_format = "#,##0"
        ws.cell(row=r, column=6, value=f["years"]).number_format = "0.0"
        ws.cell(row=r, column=7, value=rec).number_format = "0.00"
        ws.cell(row=r, column=8, value=note)

        if f["n"] < 800:
            ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=WARN)
        r += 1

    last = r - 1
    body(ws, 5, last, len(headers))
    ws.cell(row=5, column=1).font = Font(name=FONT, size=10, bold=True)

    ws.cell(row=last + 2, column=1,
            value="Sarı hücre: örneklem 800'ün altında, belirsizlik yüksek.")
    ws.cell(row=last + 2, column=1).font = Font(name=FONT, size=9, italic=True, color="92400E")

    ws.cell(row=last + 3, column=1, value="Ortalama b")
    ws.cell(row=last + 3, column=2, value=f"=AVERAGE(B6:B{last})").number_format = "0.000"
    ws.cell(row=last + 4, column=1, value="Standart sapma")
    ws.cell(row=last + 4, column=2, value=f"=STDEV(B6:B{last})").number_format = "0.000"
    ws.cell(row=last + 5, column=1, value="Toplam olay")
    ws.cell(row=last + 5, column=2, value=f"=SUM(E6:E{last})").number_format = "#,##0"
    for i in range(3, 6):
        ws.cell(row=last + i, column=1).font = Font(name=FONT, size=10, bold=True)

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Bölgelere göre b değeri"
    chart.y_axis.title = "b"
    chart.height, chart.width = 16, 22
    data = Reference(ws, min_col=2, min_row=4, max_row=last)
    cats = Reference(ws, min_col=1, min_row=5, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.legend = None
    ws.add_chart(chart, f"J4")

    return items


# ---------------------------------------------------------------- sayfa 3
def sheet_egri(wb, gr):
    ws = wb.create_sheet("GR Eğrisi")
    widths(ws, {"A": 12, "B": 16, "C": 16})

    title(ws, "Gutenberg-Richter Uyum Eğrisi (Dünya)",
          "Gözlenen birikimli sayım ve doğrusal uyum")

    g = gr.get("global")
    if not g:
        return

    for i, h in enumerate(["Büyüklük", "log₁₀(N) gözlenen", "log₁₀(N) uyum"], 1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, 3)

    r = 5
    for p in g["curve"]:
        ws.cell(row=r, column=1, value=p["mag"]).number_format = "0.0"
        ws.cell(row=r, column=2, value=p["logN"]).number_format = "0.000"
        ws.cell(row=r, column=3, value=p["fitted"]).number_format = "0.000"
        r += 1

    last = r - 1
    body(ws, 5, last, 3)

    chart = ScatterChart()
    chart.title = "log₁₀(N) = a − b·M"
    chart.x_axis.title = "Büyüklük (M)"
    chart.y_axis.title = "log₁₀(N)"
    chart.height, chart.width = 11, 20
    chart.style = 13

    xref = Reference(ws, min_col=1, min_row=5, max_row=last)
    for col, dashed in ((2, False), (3, True)):
        yref = Reference(ws, min_col=col, min_row=4, max_row=last)
        s = Series(yref, xref, title_from_data=True)
        if dashed:
            s.marker.symbol = "none"
        else:
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
        chart.series.append(s)

    ws.add_chart(chart, "E4")

    ws.cell(row=last + 2, column=1, value=f"b = {g['b']:.3f}   R² = {g['r2']:.4f}   n = {g['n']:,}")
    ws.cell(row=last + 2, column=1).font = Font(name=FONT, size=10, bold=True, color=ACCENT)


# ---------------------------------------------------------------- sayfa 4
def sheet_migration(wb, mig):
    ws = wb.create_sheet("Göç Analizi")
    widths(ws, {"A": 26, "B": 9, "C": 8, "D": 13, "E": 9, "F": 10, "G": 10,
                "H": 11, "I": 13, "J": 26})

    title(ws, "Fay Hatlarında Göç Analizi",
          "Depremler fay eksenine izdüşürülüp zamana göre regresyona sokuldu")

    headers = ["Fay hattı", "Olay", "M eşik", "Hız (km/yıl)", "Yön", "R²",
               "r", "p değeri", "Bonferroni", "Sonuç"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(headers))

    items = sorted(mig.values(), key=lambda x: x["significance"]["p"])

    r = 5
    for f in items:
        m, s = f["migration"], f["significance"]
        survives = s.get("survives_correction", False)
        result = ("Anlamlı" if survives
                  else "Ham eşiği geçti" if s["p"] < 0.05
                  else "Tesadüften ayırt edilemiyor")

        ws.cell(row=r, column=1, value=f["label"])
        ws.cell(row=r, column=2, value=f["n"]).number_format = "#,##0"
        ws.cell(row=r, column=3, value=f["min_magnitude"]).number_format = "0.0"
        ws.cell(row=r, column=4, value=m["velocity_km_per_year"]).number_format = "+0.00;-0.00"
        ws.cell(row=r, column=5, value=m["direction"])
        ws.cell(row=r, column=6, value=m["r2"]).number_format = "0.000"
        ws.cell(row=r, column=7, value=s["r"]).number_format = "+0.000;-0.000"
        ws.cell(row=r, column=8, value=s["p"]).number_format = "0.0000"
        ws.cell(row=r, column=9, value="Geçti" if survives else "Geçemedi")
        ws.cell(row=r, column=10, value=result)

        if survives:
            fill = "DCFCE7"
        elif s["p"] < 0.05:
            fill = WARN
        else:
            fill = None
        if fill:
            for c in (8, 9, 10):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=fill)
        r += 1

    last = r - 1
    body(ws, 5, last, len(headers), zebra=False)

    ws.cell(row=last + 2, column=1, value="Test edilen fay")
    ws.cell(row=last + 2, column=2, value=f"=COUNTA(A5:A{last})")
    ws.cell(row=last + 3, column=1, value="Ham eşiği geçen (p<0,05)")
    ws.cell(row=last + 3, column=2, value=f'=COUNTIF(H5:H{last},"<0.05")')
    ws.cell(row=last + 4, column=1, value="Şansla beklenen yanlış pozitif")
    ws.cell(row=last + 4, column=2, value=f"=B{last + 2}*0.05").number_format = "0.00"
    ws.cell(row=last + 5, column=1, value="Bonferroni sonrası kalan")
    ws.cell(row=last + 5, column=2, value=f'=COUNTIF(I5:I{last},"Geçti")')
    for i in range(2, 6):
        ws.cell(row=last + i, column=1).font = Font(name=FONT, size=10, bold=True)

    chart = ScatterChart()
    chart.title = "R² ve p değeri dağılımı"
    chart.x_axis.title = "p değeri"
    chart.y_axis.title = "R²"
    chart.height, chart.width = 11, 18
    xref = Reference(ws, min_col=8, min_row=5, max_row=last)
    yref = Reference(ws, min_col=6, min_row=4, max_row=last)
    s = Series(yref, xref, title_from_data=True)
    s.marker.symbol = "circle"
    s.graphicalProperties.line.noFill = True
    chart.series.append(s)
    chart.legend = None
    ws.add_chart(chart, "L4")


# ---------------------------------------------------------------- sayfa 5
def sheet_katalog(wb, decades):
    ws = wb.create_sheet("Katalog Dağılımı")
    widths(ws, {"A": 12, "B": 14, "C": 14, "D": 14})

    title(ws, "Katalog Bütünlüğü",
          "Onyıllara göre kayıt sayısı — erken dönemdeki düşüklük gerçek değil, tespit eksikliği")

    for i, h in enumerate(["Onyıl", "Toplam", "M6+", "M7+"], 1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, 4)

    r = 5
    for dec in sorted(decades):
        d = decades[dec]
        ws.cell(row=r, column=1, value=f"{dec}s")
        ws.cell(row=r, column=2, value=d["total"]).number_format = "#,##0"
        ws.cell(row=r, column=3, value=d["m6"]).number_format = "#,##0"
        ws.cell(row=r, column=4, value=d["m7"]).number_format = "#,##0"
        r += 1

    last = r - 1
    body(ws, 5, last, 4)

    chart = LineChart()
    chart.title = "Onyıllara göre kayıt sayısı"
    chart.y_axis.title = "Deprem sayısı"
    chart.x_axis.title = "Onyıl"
    chart.height, chart.width = 11, 20
    data = Reference(ws, min_col=2, max_col=4, min_row=4, max_row=last)
    cats = Reference(ws, min_col=1, min_row=5, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "F4")

    ws.cell(row=last + 2, column=1,
            value="M7+ sayısı onyıllar boyunca sabit kalır; toplam sayının artışı "
                  "sismik aktivite artışı değil, sensör ağının genişlemesidir.")
    ws.cell(row=last + 2, column=1).font = Font(name=FONT, size=9, italic=True, color="64748B")


# ---------------------------------------------------------------- sayfa 6
def sheet_yontem(wb):
    ws = wb.create_sheet("Yöntem")
    widths(ws, {"A": 30, "B": 90})
    title(ws, "Yöntem ve Varsayımlar")

    blocks = [
        ("Gutenberg-Richter", ""),
        ("Formül", "log₁₀(N) = a − b·M. N, büyüklüğü M ve üstünde olan deprem sayısı."),
        ("Uygulama", "M4.5–8.5 arası 0,5 birim aralıklarla birikimli sayım, en küçük kareler."),
        ("Bölge ataması", "Deprem koordinatları ülke poligonlarıyla eşleştirildi; "
                          "denizdekiler 3 derece içindeki en yakın kıyıya atandı."),
        ("Eşik", "300 olaydan az kayda sahip bölgeler dışarıda bırakıldı."),
        ("", ""),
        ("Göç analizi", ""),
        ("İzdüşüm", "Her deprem fay polyline'ı üzerine dik olarak izdüşürüldü; "
                    "eksenden belirlenen yarı genişlikten uzaktakiler elendi."),
        ("Regresyon", "Yıl (x) ile fay boyunca mesafe (y) arasında doğrusal uyum. "
                      "Eğimin kendisi km/yıl cinsinden göç hızıdır."),
        ("Permütasyon testi", "Depremlerin tarihleri 5.000 kez karıştırıldı. p değeri, "
                              "karıştırılmış korelasyonun gerçek korelasyondan güçlü "
                              "çıkma oranıdır."),
        ("Bonferroni", "27 bağımsız test yapıldığı için eşik 0,05/27 = 0,00185'e çekildi. "
                       "Düzeltmesiz 0,05 kullanılsaydı, tamamen rastgele veride bile "
                       "ortalama 1,35 yanlış pozitif beklenirdi."),
        ("", ""),
        ("Kritik uyarılar", ""),
        ("Büyük örneklem yanılgısı", "n büyüdükçe çok küçük korelasyonlar bile düşük p "
                                     "değeri verir. R² değeri düşükse (örn. 0,016) "
                                     "ilişki pratikte anlamsızdır."),
        ("Nedensellik", "Zamansal ve mekânsal yakınlık nedensellik göstermez. Zincirleme "
                        "analizinde gözlenen sayı, o bölgenin temel oranıyla karşılaştırılır."),
        ("Tahmin", "Bu çalışma deprem tahmini yapmaz. Depremlerin önceden tahmin edilmesi "
                   "bilimsel olarak mümkün değildir; USGS bunu açıkça belirtir."),
    ]

    r = 4
    for label, text in blocks:
        is_section = label and not text
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=text)
        ws.cell(row=r, column=1).font = Font(
            name=FONT, size=10, bold=True, color=ACCENT if is_section else "000000"
        )
        ws.cell(row=r, column=2).font = Font(name=FONT, size=10)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        if is_section:
            for c in (1, 2):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=LIGHT)
        if text and len(text) > 90:
            ws.row_dimensions[r].height = 30
        r += 1


# ---------------------------------------------------------------- ana akis
def catalog_summary(rows):
    decades = {}
    first = last = None

    for e in rows:
        year = 1970 + e["time"] / (365.25 * 86_400_000)
        first = year if first is None or year < first else first
        last = year if last is None or year > last else last

        dec = int(year // 10 * 10)
        d = decades.setdefault(dec, {"total": 0, "m6": 0, "m7": 0})
        d["total"] += 1
        if e["mag"] >= 6:
            d["m6"] += 1
        if e["mag"] >= 7:
            d["m7"] += 1

    return {
        "total": len(rows),
        "first_year": int(first),
        "last_year": int(last),
    }, decades


def main():
    gr = json.loads(GUTENBERG.read_text(encoding="utf-8"))
    mig = json.loads(MIGRATION.read_text(encoding="utf-8"))
    rows = json.loads(CATALOG.read_text(encoding="utf-8"))

    stats, decades = catalog_summary(rows)

    wb = Workbook()
    sheet_ozet(wb, stats, len([k for k in gr if k != "global"]), len(mig))
    sheet_gutenberg(wb, gr)
    sheet_egri(wb, gr)
    sheet_migration(wb, mig)
    sheet_katalog(wb, decades)
    sheet_yontem(wb)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    print(f"Rapor olusturuldu: {OUT}")
    print(f"  Sayfa sayisi : {len(wb.sheetnames)}")
    print(f"  Sayfalar     : {', '.join(wb.sheetnames)}")
    print(f"  Boyut        : {OUT.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main()